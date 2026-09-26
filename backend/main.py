from __future__ import annotations 

import asyncio 
import csv 
import io 
import json 
import logging 
import math 
import os 
import re 
import time 
import uuid 
import hashlib 
from collections import Counter ,defaultdict 
from contextlib import asynccontextmanager 
from datetime import datetime ,timezone 
from pathlib import Path 
from threading import Lock 
from typing import Any ,Dict ,List ,Optional ,Tuple 

import httpx 
from dotenv import load_dotenv 
from fastapi import FastAPI ,File ,HTTPException ,Query ,Request ,UploadFile 
from fastapi .concurrency import run_in_threadpool 
from fastapi .middleware .cors import CORSMiddleware 
from fastapi .responses import JSONResponse 
from pydantic import BaseModel ,Field 
from pypdf import PdfReader 

try :
    from docx import Document as DocxDocument 
except ImportError :
    DocxDocument =None 

try :
    from supabase import Client ,create_client 
except ImportError :
    Client =Any 
    create_client =None 

try :
    from PIL import Image 
except ImportError :
    Image =None 

try :
    import spacy 
except ImportError :
    spacy =None 

try :
    import fitz 
except ImportError :
    fitz =None 


try :
    from openpyxl import Workbook 
    from openpyxl .styles import Font ,PatternFill ,Alignment ,Border ,Side 
    from openpyxl .utils import get_column_letter 
except ImportError :
    Workbook =None 
    Font =PatternFill =Alignment =Border =Side =get_column_letter =None 

try :
    import pytesseract 
    try :

        if os .name =="nt":
            common_tesseract_paths =[
            r"C:\\Program Files\\Tesseract-OCR\\tesseract.exe",
            r"C:\\Program Files (x86)\\Tesseract-OCR\\tesseract.exe",
            os .path .expandvars (r"%LOCALAPPDATA%\\Tesseract-OCR\\tesseract.exe"),
            ]
            for _tesseract_path in common_tesseract_paths :
                if os .path .exists (_tesseract_path ):
                    pytesseract .pytesseract .tesseract_cmd =_tesseract_path 
                    break 
    except Exception :
        pass 
except ImportError :
    pytesseract =None 






load_dotenv ()

APP_NAME =os .getenv ("APP_NAME","DocuSense AI")
APP_VERSION =os .getenv ("APP_VERSION","4.0.0")
ENVIRONMENT =os .getenv ("ENVIRONMENT","development")


MAX_INDEX_SENTENCES =int (os .getenv ("MAX_INDEX_SENTENCES","20000"))
MAX_INDEX_CHARS =int (os .getenv ("MAX_INDEX_CHARS","2500000"))
MAX_CONTEXT_CHARS =int (os .getenv ("MAX_CONTEXT_CHARS","18000"))
AI_CACHE_TTL_SECONDS =int (os .getenv ("AI_CACHE_TTL_SECONDS","300"))
ADVANCED_ENTITY_LIMIT =int (os .getenv ("ADVANCED_ENTITY_LIMIT","100"))

AI_RESPONSE_CACHE :Dict [str ,Tuple [float ,Dict [str ,Any ]]]={}
AI_RESPONSE_CACHE_LOCK =Lock ()
DOCUMENT_FINGERPRINTS :Dict [str ,str ]={}
DOCUMENT_FINGERPRINT_LOCK =Lock ()

UPLOAD_DIR =Path (os .getenv ("UPLOAD_DIR","uploads"))
UPLOAD_DIR .mkdir (parents =True ,exist_ok =True )

MAX_FILE_SIZE_MB =int (os .getenv ("MAX_FILE_SIZE_MB","25"))
MAX_FILE_SIZE =MAX_FILE_SIZE_MB *1024 *1024 
MAX_TEXT_CHARS =int (os .getenv ("MAX_TEXT_CHARS","1000000"))
MAX_CHAT_CONTEXT_CHARS =int (os .getenv ("MAX_CHAT_CONTEXT_CHARS","24000"))
SEARCH_INDEX_LIMIT =int (os .getenv ("SEARCH_INDEX_LIMIT","12000"))
BATCH_CONCURRENCY =int (os .getenv ("BATCH_CONCURRENCY","4"))
AI_CONNECT_TIMEOUT =float (os .getenv ("AI_CONNECT_TIMEOUT","5"))
AI_READ_TIMEOUT =float (os .getenv ("AI_READ_TIMEOUT","25"))
MAX_BATCH_FILES =int (os .getenv ("MAX_BATCH_FILES","20"))
REQUEST_TIMEOUT =int (os .getenv ("REQUEST_TIMEOUT","60"))
FILE_RETENTION_HOURS =int (os .getenv ("FILE_RETENTION_HOURS","24"))
RATE_LIMIT_PER_MINUTE =int (os .getenv ("RATE_LIMIT_PER_MINUTE","120"))

ALLOWED_EXTENSIONS ={
".pdf",".docx",".txt",".md",".csv",
".png",".jpg",".jpeg",".webp",
}

CORS_ORIGINS_RAW =os .getenv (
"CORS_ORIGINS",
"http://localhost:5173,http://127.0.0.1:5173",
)
CORS_ORIGINS =[
origin .strip ()for origin in CORS_ORIGINS_RAW .split (",")if origin .strip ()
]

SUPABASE_URL =os .getenv ("SUPABASE_URL")
SUPABASE_KEY =os .getenv ("SUPABASE_KEY")

RAPIDAPI_KEY =os .getenv ("RAPIDAPI_KEY")
RAPIDAPI_HOST =os .getenv ("RAPIDAPI_HOST")
RAPIDAPI_ENDPOINT =os .getenv ("RAPIDAPI_ENDPOINT")


TESSERACT_CMD =os .getenv ("TESSERACT_CMD")
if TESSERACT_CMD and pytesseract :
    pytesseract .pytesseract .tesseract_cmd =TESSERACT_CMD 

logger =logging .getLogger ("docusense")
logging .basicConfig (
level =os .getenv ("LOG_LEVEL","INFO").upper (),
format ="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)






supabase :Optional [Client ]=None 
LOCAL_DOCUMENTS :Dict [str ,Dict [str ,Any ]]={}
LOCAL_LOCK =Lock ()




DOCUMENT_CACHE :Dict [str ,Dict [str ,Any ]]={}
DOCUMENT_CACHE_LOCK =Lock ()

def cache_document (document_id :str ,record :Dict [str ,Any ])->None :
    analysis =record .get ("analysis")or {}
    page_data =analysis .get ("page_data")or []
    text ="\n".join (
    item .get ("text","")for item in page_data if isinstance (item ,dict )
    )
    index =build_search_index (text )if text else []
    with DOCUMENT_CACHE_LOCK :
        DOCUMENT_CACHE [document_id ]={
        "record":record ,
        "text":text ,
        "search_index":index ,
        "cached_at":time .time (),
        }

def get_cached_document (document_id :str )->Optional [Dict [str ,Any ]]:
    with DOCUMENT_CACHE_LOCK :
        item =DOCUMENT_CACHE .get (document_id )
        return item .copy ()if item else None 

if SUPABASE_URL and SUPABASE_KEY and create_client :
    try :
        supabase =create_client (SUPABASE_URL ,SUPABASE_KEY )
        logger .info ("Supabase persistence enabled.")
    except Exception as exc :
        logger .warning ("Supabase initialization failed: %s",exc )
        supabase =None 
else :
    logger .info ("Supabase not configured; enhanced local persistence is enabled.")


def utc_now ()->str :
    return datetime .now (timezone .utc ).isoformat ()


def persistence_mode ()->str :
    return "supabase+local-cache"if supabase else "local-memory"


def ai_provider_available ()->bool :
    return bool (RAPIDAPI_KEY and RAPIDAPI_HOST and RAPIDAPI_ENDPOINT )






@asynccontextmanager 
async def lifespan (_app :FastAPI ):
    removed =cleanup_old_uploads ()
    logger .info (
    "%s %s started. Removed %s expired upload(s).",
    APP_NAME ,APP_VERSION ,removed ,
    )
    yield 
    logger .info ("%s shutting down.",APP_NAME )


app =FastAPI (
title =APP_NAME ,
description ="Production-oriented AI document intelligence API.",
version =APP_VERSION ,
lifespan =lifespan ,
)

app .add_middleware (
CORSMiddleware ,
allow_origins =CORS_ORIGINS ,
allow_credentials =True ,
allow_methods =["*"],
allow_headers =["*"],
)

_RATE_BUCKETS :Dict [str ,List [float ]]=defaultdict (list )
_RATE_LOCK =Lock ()


@app .middleware ("http")
async def request_middleware (request :Request ,call_next ):
    started =time .perf_counter ()
    request_id =request .headers .get ("X-Request-ID")or uuid .uuid4 ().hex 

    client_host =request .client .host if request .client else "unknown"
    now =time .time ()
    with _RATE_LOCK :
        bucket =[t for t in _RATE_BUCKETS [client_host ]if now -t <60 ]
        bucket .append (now )
        _RATE_BUCKETS [client_host ]=bucket 
        limited =len (bucket )>RATE_LIMIT_PER_MINUTE 

    if limited :
        return JSONResponse (
        status_code =429 ,
        content ={
        "status":"error",
        "detail":"Rate limit exceeded. Please retry shortly.",
        "request_id":request_id ,
        },
        headers ={"Retry-After":"60","X-Request-ID":request_id },
        )

    try :
        response =await call_next (request )
    except Exception :
        logger .exception ("Unhandled request error | request_id=%s",request_id )
        raise 
    finally :
        elapsed =(time .perf_counter ()-started )*1000 
        logger .info (
        "%s %s -> completed in %.2fms | request_id=%s",
        request .method ,request .url .path ,elapsed ,request_id ,
        )

    response .headers ["X-Request-ID"]=request_id 
    response .headers ["X-Process-Time-Ms"]=f"{elapsed :.2f}"
    return response 






class ChatRequest (BaseModel ):
    question :str =Field (...,min_length =1 ,max_length =5000 )
    document_id :Optional [str ]=None 
    context :Optional [str ]=Field (default =None ,max_length =MAX_CHAT_CONTEXT_CHARS )


class CompareRequest (BaseModel ):
    document_ids :List [str ]=Field (...,min_length =2 ,max_length =10 )


class SearchRequest (BaseModel ):
    query :str =Field (...,min_length =1 ,max_length =500 )
    document_id :Optional [str ]=None 
    limit :int =Field (default =10 ,ge =1 ,le =50 )


class HealthResponse (BaseModel ):
    status :str 
    service :str 
    version :str 
    environment :str 
    persistence :str 
    ai_provider :str 






RISK_KEYWORDS :Dict [str ,List [str ]]={
"Critical":[
"confession of judgment",
"unlimited liability",
"personal guarantee",
"waive all claims",
"irrevocable",
"without limitation",
],
"High":[
"liquidated damages",
"indemnify",
"indemnification",
"sole discretion",
"uncapped liability",
"exclusive remedy",
"non-refundable",
"termination fee",
],
"Medium":[
"non-compete",
"non compete",
"penalty",
"automatic renewal",
"auto-renewal",
"exclusivity",
"termination for convenience",
"late fee",
"interest rate",
"arbitration",
"jurisdiction",
"confidentiality",
"non-solicitation",
"notice period",
],
"Low":[
"payment terms",
"governing law",
"renewal",
"termination",
"notice",
"confidential",
],
}

TAG_KEYWORDS :Dict [str ,List [str ]]={
"NDA":["non-disclosure","nondisclosure","confidential information","confidentiality agreement","nda"],
"Employment":["employee","employer","employment","salary","compensation","job title","benefits"],
"Invoice":["invoice","amount due","subtotal","tax","balance due","payment terms"],
"Contract":["agreement","party","parties","hereby","obligations","terms and conditions"],
"Legal":["governing law","jurisdiction","arbitration","dispute resolution","legal notice"],
"Resume":["curriculum vitae","resume","work experience","professional summary","education","skills:"],
"Financial":["revenue","profit","balance sheet","cash flow","financial statement","accounts payable","accounts receivable"],
"Policy":["policy","procedure","compliance","employees must","must comply"],
"Report":["executive summary","findings","recommendations","methodology","conclusion"],
"Purchase Order":["purchase order","po number","ship to","bill to"],
}

DOCUMENT_TYPE_DESCRIPTIONS ={
"NDA":"A non-disclosure agreement covering confidential information and disclosure restrictions.",
"Employment":"An employment document covering compensation, responsibilities, benefits, or termination.",
"Invoice":"A billing document containing charges, taxes, payment information, or amounts due.",
"Contract":"An agreement defining obligations, rights, terms, and conditions between parties.",
"Legal":"A document containing legal, jurisdictional, arbitration, or dispute-resolution provisions.",
"Resume":"A résumé or CV summarizing professional experience, education, and skills.",
"Financial":"A financial document containing accounting or monetary performance information.",
"Policy":"An organizational policy or procedure defining rules, requirements, or compliance expectations.",
"Report":"A structured report containing findings, analysis, methodology, or recommendations.",
"Purchase Order":"A purchasing document describing ordered goods/services and commercial terms.",
}

DEFAULT_CATEGORY ="General document"
DEFAULT_CATEGORY_DESCRIPTION ="No specific document type was confidently identified."

DATE_PATTERN =re .compile (
r"\b(?:January|February|March|April|May|June|July|August|September|"
r"October|November|December)\s+\d{1,2}(?:st|nd|rd|th)?(?:,)?\s+\d{4}\b",
re .IGNORECASE ,
)
NUMERIC_DATE_PATTERN =re .compile (
r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2})\b"
)
MONEY_PATTERN =re .compile (
r"(?:[$€£₨]\s?\d[\d,]*(?:\.\d{2})?|\b\d[\d,]*(?:\.\d{2})?\s?(?:USD|EUR|GBP|PKR)\b)",
re .IGNORECASE ,
)
EMAIL_PATTERN =re .compile (r"\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}\b")
PHONE_PATTERN =re .compile (
r"(?<!\d)(?:\+\d{1,3}[\s.-]?)?(?:\(?\d{2,4}\)?[\s.-]?)?"
r"\d{3,4}[\s.-]?\d{3,4}(?!\d)"
)
URL_PATTERN =re .compile (r"\bhttps?://[^\s<>\"]+",re .IGNORECASE )
SECTION_PATTERN =re .compile (
r"(?im)^(?:\d+(?:\.\d+)*[.)]?\s+)?"
r"(?:[A-Z][A-Z0-9 &,/()'_-]{3,}|"
r"(?:ARTICLE|SECTION|CLAUSE|SCHEDULE|EXHIBIT)\s+[A-Z0-9.-]+)\s*$"
)

POSITIVE_WORDS ={
"benefit","success","successful","improve","improved","positive",
"excellent","strong","achieve","growth","secure","approved",
}
NEGATIVE_WORDS ={
"risk","loss","failure","penalty","breach","terminate","termination",
"liability","dispute","warning","late","damage","uncertain",
}

ACTION_PATTERNS =[
r"\b(?:must|shall|required to|needs to|need to)\b[^.!?\n]{5,180}",
r"\b(?:submit|provide|review|sign|pay|renew|terminate|notify|deliver|respond|approve)\b[^.!?\n]{5,180}",
]

DEADLINE_PATTERN =re .compile (
r"\b(?:due|deadline|expires?|expiration|by|before|within)\b[^.!?\n]{0,100}"
r"(?:\d+\s+(?:days?|weeks?|months?|years?)|"
r"(?:January|February|March|April|May|June|July|August|September|October|November|December)"
r"\s+\d{1,2}(?:st|nd|rd|th)?(?:,)?\s+\d{4})",
re .IGNORECASE ,
)






def safe_filename (filename :Optional [str ])->str :
    original =Path (filename or "document").name 
    cleaned =re .sub (r"[^A-Za-z0-9._-]+","_",original ).strip ("._")
    return cleaned [:180 ]or "document"


def make_document_id (filename :str )->str :
    stem =Path (filename ).stem .lower ()
    stem =re .sub (r"[^a-z0-9]+","-",stem ).strip ("-")
    return f"{stem or 'document'}-{uuid .uuid4 ().hex [:10 ]}"


def normalize_text (text :str )->str :
    text =text .replace ("\x00"," ")
    text =text .replace ("\r\n","\n").replace ("\r","\n")
    text =re .sub (r"[ \t]+"," ",text )
    text =re .sub (r"\n{3,}","\n\n",text )
    return text .strip ()


def truncate_text (text :str ,limit :int =MAX_TEXT_CHARS )->str :
    if len (text )<=limit :
        return text 
    return text [:limit ]+"\n\n[Text truncated by server limit.]"


def word_count (text :str )->int :
    return len (re .findall (r"\b[\w'-]+\b",text ,flags =re .UNICODE ))


def sentence_split (text :str )->List [str ]:
    return [
    s .strip ()
    for s in re .split (r"(?<=[.!?])\s+",text )
    if s .strip ()
    ]


def unique_preserve (values :List [str ],limit :int =50 )->List [str ]:
    seen =set ()
    result =[]
    for value in values :
        normalized =str (value ).strip ()
        if not normalized :
            continue 
        key =normalized .lower ()
        if key in seen :
            continue 
        seen .add (key )
        result .append (normalized )
        if len (result )>=limit :
            break 
    return result 


def safe_json (data :Any )->Any :
    try :
        json .dumps (data )
        return data 
    except TypeError :
        return json .loads (json .dumps (data ,default =str ))


def detect_language (text :str )->str :
    if not text .strip ():
        return "Unknown"
    sample =text [:10000 ]
    latin =len (re .findall (r"[A-Za-z]",sample ))
    english =len (re .findall (
    r"\b(?:the|and|of|to|in|for|with|is|this|that|from|are)\b",
    sample ,re .I ,
    ))
    if latin and english >=2 :
        return "English"
    if re .search (r"[\u0600-\u06FF]",sample ):
        return "Arabic/Urdu"
    if re .search (r"[\u4E00-\u9FFF]",sample ):
        return "Chinese"
    if re .search (r"[\u3040-\u30FF]",sample ):
        return "Japanese"
    return "Unknown"


def readability_score (text :str )->int :
    words =word_count (text )
    sentences =max (1 ,len (sentence_split (text )))
    if words <20 :
        return 0 
    syllables =0 
    for word in re .findall (r"[A-Za-z]+",text ):
        groups =re .findall (r"[aeiouy]+",word .lower ())
        syllables +=max (1 ,len (groups ))
    score =206.835 -1.015 *(words /sentences )-84.6 *(syllables /max (1 ,words ))
    return max (0 ,min (100 ,round (score )))


def text_sentiment (text :str )->Dict [str ,Any ]:
    words =re .findall (r"[A-Za-z']+",text .lower ())
    positive =sum (w in POSITIVE_WORDS for w in words )
    negative =sum (w in NEGATIVE_WORDS for w in words )
    total =max (1 ,positive +negative )
    if positive >negative *1.25 :
        tone ="Positive"
    elif negative >positive *1.25 :
        tone ="Cautionary"
    else :
        tone ="Neutral"
    return {
    "tone":tone ,
    "positive_signals":positive ,
    "negative_signals":negative ,
    "signal_balance":round ((positive -negative )/total ,3 ),
    }






def extract_pdf_text(file_path: Path) -> Tuple[int, str, List[Dict[str, Any]]]:
    try:
        reader = PdfReader(str(file_path))
    except Exception as exc:
        raise HTTPException(400, f"Could not read PDF: {exc}")
    pages = []
    chunks = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = normalize_text(page.extract_text() or "")
        except Exception as exc:
            logger.warning("PDF page %s extraction failed: %s", page_number, exc)
            page_text = ""
        used_ocr = False
        if not page_text.strip() and fitz is not None and pytesseract is not None and Image is not None:
            try:
                pdf_doc = fitz.open(str(file_path))
                pix = pdf_doc[page_number - 1].get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                page_text = normalize_text(pytesseract.image_to_string(image))
                used_ocr = bool(page_text.strip())
                pdf_doc.close()
            except Exception as exc:
                logger.warning("PDF OCR page %s failed: %s", page_number, exc)
        chunks.append(page_text)
        pages.append({"page": page_number, "characters": len(page_text), "words": word_count(page_text), "text": page_text, "ocr": used_ocr})
    return len(reader.pages), truncate_text("\n\n".join(chunks)), pages


def extract_docx_text (file_path :Path )->Tuple [int ,str ,List [Dict [str ,Any ]]]:
    if DocxDocument is None :
        raise HTTPException (500 ,"DOCX support unavailable. Install python-docx.")
    try :
        document =DocxDocument (str (file_path ))
    except Exception as exc :
        raise HTTPException (400 ,f"Could not read DOCX: {exc }")

    paragraphs =[
    normalize_text (p .text )
    for p in document .paragraphs 
    if p .text .strip ()
    ]
    table_text =[]
    for table_index ,table in enumerate (document .tables ,start =1 ):
        table_text .append (f"[Table {table_index }]")
        for row in table .rows :
            cells =[normalize_text (cell .text )for cell in row .cells ]
            table_text .append (" | ".join (cells ))

    combined =truncate_text (normalize_text ("\n".join (paragraphs +table_text )))
    return 1 ,combined ,[{
    "page":1 ,
    "characters":len (combined ),
    "words":word_count (combined ),
    "text":combined ,
    "tables":len (document .tables ),
    }]


def extract_plain_text (file_path :Path )->Tuple [int ,str ,List [Dict [str ,Any ]]]:
    try :
        text =file_path .read_text (encoding ="utf-8",errors ="replace")
    except Exception as exc :
        raise HTTPException (400 ,f"Could not read text file: {exc }")
    text =truncate_text (normalize_text (text ))
    return 1 ,text ,[{
    "page":1 ,
    "characters":len (text ),
    "words":word_count (text ),
    "text":text ,
    }]


def extract_csv_text (file_path :Path )->Tuple [int ,str ,List [Dict [str ,Any ]]]:
    try :
        raw =file_path .read_text (encoding ="utf-8",errors ="replace")
        rows =list (csv .reader (io .StringIO (raw )))
    except Exception as exc :
        raise HTTPException (400 ,f"Could not read CSV: {exc }")

    lines =[" | ".join (cell .strip ()for cell in row )for row in rows ]
    text =truncate_text (normalize_text ("\n".join (lines )))
    return 1 ,text ,[{
    "page":1 ,
    "characters":len (text ),
    "words":word_count (text ),
    "rows":len (rows ),
    "columns":max ((len (r )for r in rows ),default =0 ),
    "text":text ,
    }]


def extract_image_text (file_path :Path )->Tuple [int ,str ,List [Dict [str ,Any ]]]:
    if Image is None :
        raise HTTPException (
        501 ,
        "Image OCR requires Pillow. Install it with: pip install pillow",
        )

    if pytesseract is None :
        raise HTTPException (
        501 ,
        "Image OCR is unavailable because pytesseract is not installed. "
        "Install it with: pip install pytesseract",
        )

    try :


        pytesseract .get_tesseract_version ()
    except Exception :
        raise HTTPException (
        501 ,
        "pytesseract is installed, but the Tesseract-OCR program was not found. "
        "Install Tesseract OCR for Windows and restart the server. "
        "If it is installed in a custom location, set TESSERACT_CMD in .env.",
        )

    try :
        image =Image .open (file_path )
        text =pytesseract .image_to_string (image )
    except Exception as exc :
        raise HTTPException (400 ,f"Could not OCR image: {exc }")

    text =truncate_text (normalize_text (text ))
    return 1 ,text ,[{
    "page":1 ,
    "characters":len (text ),
    "words":word_count (text ),
    "text":text ,
    "ocr":True ,
    }]


def extract_document (file_path :Path )->Tuple [int ,str ,List [Dict [str ,Any ]]]:
    extension =file_path .suffix .lower ()
    if extension ==".pdf":
        return extract_pdf_text (file_path )
    if extension ==".docx":
        return extract_docx_text (file_path )
    if extension in {".txt",".md"}:
        return extract_plain_text (file_path )
    if extension ==".csv":
        return extract_csv_text (file_path )
    if extension in {".png",".jpg",".jpeg",".webp"}:
        return extract_image_text (file_path )
    raise HTTPException (415 ,f"Unsupported file type: {extension or 'unknown'}")






def find_keyword_context (text :str ,keyword :str ,radius :int =220 )->str :
    match =re .search (re .escape (keyword ),text ,re .IGNORECASE )
    if not match :
        return ""
    start =max (0 ,match .start ()-radius )
    end =min (len (text ),match .end ()+radius )
    return re .sub (r"\s+"," ",text [start :end ]).strip ()


def analyze_risk (text :str )->Tuple [str ,List [Dict [str ,Any ]]]:
    normalized =re .sub (r"\s+"," ",text or "").strip ()
    lower =normalized .lower ()
    findings =[]
    score =0 
    weights ={"Critical":45 ,"High":22 ,"Medium":9 ,"Low":2 }
    patterns ={
    "Critical":[
    "unlimited liability","uncapped liability","unlimited damages","personal guarantee",
    "personally guarantees","irrevocably guarantees","confession of judgment",
    "waive all claims","waives all claims","waive all defenses","waives all defenses",
    "waives the right to sue","sole discretion","without limitation of liability",
    "all losses damages penalties costs","personally liable","joint and several liability",
    "no liability cap","unrestricted access to personal data"
    ],
    "High":[
    "liquidated damages","indemnify and hold harmless","indemnify, defend, and hold harmless",
    "termination fee","termination penalty","automatic renewal","automatically renews",
    "non-compete","non compete","exclusivity","exclusive arbitration","binding arbitration",
    "class action waiver","jury trial waiver","unilateral","change pricing",
    "sole discretion","penalty","late payment interest","default interest",
    "4% per month","5% per month","6% per month","7% per month","8% per month",
    "9% per month","10% per month","five-year","five years","indefinitely after termination",
    "survives indefinitely","irrevocable","non-refundable","no refund","data breach notification"
    ],
    "Medium":[
    "late fee","interest rate","notice period","confidentiality","non-solicitation",
    "governing law","jurisdiction","termination","renewal","indemnification",
    "payment terms","material breach","audit rights","insurance","warranty",
    "service level","sla","processing fee","delivery deadline","retention period"
    ],
    }
    def add (level ,phrase ,count ,reason ,kind ="risk_signal"):
        nonlocal score 
        if count <=0 :
            return 
        contribution =min (weights [level ]*count ,weights [level ]*4 )
        score +=contribution 
        findings .append ({
        "level":level ,"type":kind ,"keyword":phrase ,"count":count ,
        "weight":weights [level ],"score_contribution":contribution ,
        "text":reason ,"context":find_keyword_context (normalized ,phrase ,320 )
        })
    for level ,phrases in patterns .items ():
        for phrase in phrases :
            count =len (re .findall (re .escape (phrase ),lower ,flags =re .IGNORECASE ))
            add (level ,phrase ,count ,f"{level }-risk language detected: {phrase }.")
    money =MONEY_PATTERN .findall (normalized )
    if money :
        add ("Medium","monetary obligation",min (len (money ),5 ),f"{len (money )} monetary value(s) detected; review amounts, payment obligations and penalties.")
    if re .search (r"\b(?:4|5|6|7|8|9|10|12|15|18|20)%\s*(?:per month|monthly)",lower ):
        add ("High","high monthly interest",1 ,"A potentially high recurring monthly interest rate was detected.")
    if re .search (r"\b(?:three|four|five|six|seven|eight|nine|ten|eleven|twelve|\d+)[ -]year",lower ):
        add ("High","long-term commitment",1 ,"A multi-year commitment was detected.")
    if re .search (r"\b(?:30|60|90|120|180|365)\s+days?\b",lower ):
        add ("Medium","extended notice or deadline",1 ,"A material notice or deadline period was detected.")
    missing =[]
    if "limitation of liability"not in lower and "limited liability"not in lower and "liability cap"not in lower :
        missing .append ("No clear liability cap detected.")
    if "refund"not in lower and "return of payment"not in lower :
        missing .append ("No refund or payment-return procedure detected.")
    if "governing law"not in lower and "laws of"not in lower :
        missing .append ("No governing-law provision detected.")
    if "termination"not in lower and "cancel"not in lower :
        missing .append ("No clear termination or cancellation provision detected.")
    for item in missing [:4 ]:
        add ("Medium","missing protection",1 ,item ,"missing_protection")
    critical =sum (f ["count"]for f in findings if f ["level"]=="Critical")
    high =sum (f ["count"]for f in findings if f ["level"]=="High")
    medium =sum (f ["count"]for f in findings if f ["level"]=="Medium")
    if critical :
        risk ="High"
        severity ="Critical"
    elif high >=3 or score >=65 :
        risk ="High"
        severity ="High"
    elif high >=1 or medium >=3 or score >=22 :
        risk ="Medium"
        severity ="Medium"
    else :
        risk ="Low"
        severity ="Low"
    if not findings :
        findings .append ({"level":"Low","type":"scan_result","keyword":None ,"count":0 ,"weight":0 ,"score_contribution":0 ,"text":"No configured risk signals were detected.","context":""})
    order ={"Critical":4 ,"High":3 ,"Medium":2 ,"Low":1 }
    findings .sort (key =lambda x :(-order .get (x .get ("level","Low"),1 ),-int (x .get ("count")or 0 ),x .get ("keyword")or ""))
    return risk ,findings 

def detect_tags (text :str )->Tuple [List [str ],str ,str ]:
    lower =text .lower ()
    scores ={}
    for tag ,keywords in TAG_KEYWORDS .items ():
        score =sum (1 for keyword in keywords if keyword .lower ()in lower )
        if score :
            scores [tag ]=score 

    tags =[
    tag for tag ,_score in sorted (
    scores .items (),key =lambda item :(-item [1 ],item [0 ])
    )
    ]
    category =tags [0 ]if tags else DEFAULT_CATEGORY 
    return tags ,category ,DOCUMENT_TYPE_DESCRIPTIONS .get (
    category ,DEFAULT_CATEGORY_DESCRIPTION 
    )


def extract_entities (text :str )->List [Dict [str ,Any ]]:
    entities =[]
    dates =unique_preserve (DATE_PATTERN .findall (text )+NUMERIC_DATE_PATTERN .findall (text ),50 )
    amounts =unique_preserve (MONEY_PATTERN .findall (text ),50 )
    emails =unique_preserve (EMAIL_PATTERN .findall (text ),50 )
    phones =unique_preserve (PHONE_PATTERN .findall (text ),50 )
    urls =unique_preserve (URL_PATTERN .findall (text ),50 )
    person_patterns =[
    r"\b(?:Mr|Mrs|Ms|Miss|Dr|Prof|Engr|Attorney|Adv)\.?\s+([A-Z][a-z]{1,30}(?:\s+[A-Z][a-z]{1,30}){0,3})",
    r"\b(?:name of|named|prepared by|submitted by|authorized representative|representative|employee|employer|client|customer|vendor|supplier|buyer|seller|landlord|tenant|borrower|lender)\s*[:\-]?\s*([A-Z][a-z]{1,30}(?:\s+[A-Z][a-z]{1,30}){0,3})",
    r"\bbetween\s+([A-Z][a-z]{1,30}(?:\s+[A-Z][a-z]{1,30}){0,3})\s+and\s+([A-Z][a-z]{1,30}(?:\s+[A-Z][a-z]{1,30}){0,2})(?=\s|[.,;:])",
    ]
    people =[]
    for pattern in person_patterns :
        for match in re .finditer (pattern ,text ,re .IGNORECASE ):
            values =match .groups ()
            for value in values :
                clean =re .sub (r"\s+"," ",value ).strip (" .,:;()[]{}")
                if clean and len (clean .split ())>=2 and clean .lower ()not in {"the company","the party","the employee","the employer"} and not re.search (r"\b(?:llc|ltd|limited|inc|incorporated|corporation|corp|plc)\.?$",clean,re.I):
                    people .append (clean )
    if spacy :
        try :
            nlp =spacy .load ("en_core_web_sm")
            doc =nlp (text [:200000 ])
            people .extend (ent .text .strip ()for ent in doc .ents if ent .label_ =="PERSON")
        except Exception :
            pass 
    people =unique_preserve (people ,50 )
    organizations =[]
    org_patterns = [
    r"\b(?:company|organization|corporation|corp|institution|university|bank)\s*[:\-]\s*([A-Z][A-Za-z0-9&.,' -]{2,80}?)(?=\s+(?:and|between|the|is|was|has)\b|[.;,]|$)",
    r"\b([A-Z][A-Za-z0-9&.' -]{1,30}(?:\s+[A-Z][A-Za-z0-9&.' -]{1,30}){0,5}\s+(?:LLC|LTD|LIMITED|INC|INCORPORATED|CORPORATION|CORP\.?|PLC))\b"
    ]
    for pattern in org_patterns:
        organizations.extend(m.group(1).strip(" .,:;") for m in re.finditer(pattern, text))
    if spacy :
        try :
            nlp =spacy .load ("en_core_web_sm")
            doc =nlp (text [:200000 ])
            organizations .extend (ent .text .strip ()for ent in doc .ents if ent .label_ =="ORG")
        except Exception :
            pass 
    organizations .extend (re.findall (r"\b([A-Z][A-Za-z0-9&.' -]{2,80}\s+(?:LLC|LTD|LIMITED|INC|INCORPORATED|CORPORATION|CORP\.?|PLC))\b",text ))
    organizations =unique_preserve ([x .strip (" .,:;")for x in organizations if not (re.search (r"\bbetween\b",x,re.I) and re.search (r"\band\b",x,re.I))],50 )
    for i ,value in enumerate (people ,1 ):
        entities .append ({"label":"Person"if len (people )==1 else f"Person {i }","value":value ,"type":"person"})
    for i ,value in enumerate (organizations ,1 ):
        entities .append ({"label":"Organization"if len (organizations )==1 else f"Organization {i }","value":value ,"type":"organization"})
    for i ,value in enumerate (dates ,1 ):
        entities .append ({"label":"Date"if len (dates )==1 else f"Date {i }","value":value ,"type":"date"})
    for i ,value in enumerate (amounts ,1 ):
        entities .append ({"label":"Amount"if len (amounts )==1 else f"Amount {i }","value":value ,"type":"money"})
    for i ,value in enumerate (emails ,1 ):
        entities .append ({"label":"Email"if len (emails )==1 else f"Email {i }","value":value ,"type":"email"})
    for i ,value in enumerate (phones ,1 ):
        entities .append ({"label":"Phone"if len (phones )==1 else f"Phone {i }","value":value ,"type":"phone"})
    for i ,value in enumerate (urls ,1 ):
        entities .append ({"label":"URL"if len (urls )==1 else f"URL {i }","value":value ,"type":"url"})
    return entities 

def extract_sections (text :str )->List [str ]:
    sections =[]
    for line in text .splitlines ():
        clean =re .sub (r"\s+"," ",line ).strip ()
        if 0 <len (clean )<=120 and SECTION_PATTERN .match (clean ):
            sections .append (clean )
    return unique_preserve (sections ,75 )


def extract_deadlines (text :str )->List [str ]:
    results =DEADLINE_PATTERN .findall (text )
    return unique_preserve ([re .sub (r"\s+"," ",x ).strip ()for x in results ],20 )


def extract_action_items (text :str )->List [str ]:
    results =[]
    for pattern in ACTION_PATTERNS :
        results .extend (re .findall (pattern ,text ,flags =re .IGNORECASE ))
    return unique_preserve (
    [re .sub (r"\s+"," ",item ).strip (" -:;")for item in results ],
    20 ,
    )


def extract_obligations (text :str )->List [Dict [str ,Any ]]:
    obligations =[]
    sentences =sentence_split (text )
    for sentence in sentences :
        lower =sentence .lower ()
        if any (term in lower for term in ("shall "," must ","required to","agrees to","will ")):
            if 20 <=len (sentence )<=600 :
                obligations .append ({
                "text":sentence ,
                "strength":"mandatory"if "shall"in lower or "must"in lower else "commitment",
                })
        if len (obligations )>=15 :
            break 
    return obligations 


def detect_missing_information (text :str ,category :str )->List [str ]:
    lower =text .lower ()
    expected_terms ={
    "Contract":[
    ("parties","Named parties are not obvious."),
    ("termination","Termination terms were not detected."),
    ("governing law","Governing-law language was not detected."),
    ("effective","An effective/start date was not detected."),
    ],
    "Employment":[
    ("salary","Compensation/salary language was not detected."),
    ("termination","Termination language was not detected."),
    ("benefits","Benefits language was not detected."),
    ],
    "Invoice":[
    ("amount","An amount or balance field was not detected."),
    ("due","A payment due-date/term was not detected."),
    ],
    "NDA":[
    ("confidential","Confidentiality language was not detected."),
    ("term","A term/duration provision was not detected."),
    ],
    "Purchase Order":[
    ("quantity","A quantity was not detected."),
    ("price","Pricing language was not detected."),
    ("delivery","Delivery language was not detected."),
    ],
    }
    return [
    message for term ,message in expected_terms .get (category ,[])
    if term not in lower 
    ]


def build_key_points (text :str ,findings :List [Dict [str ,Any ]])->List [str ]:
    signal_words =(
    "shall","must","may","required","payment","termination",
    "confidential","liable","liability","renewal","deadline",
    "effective","obligation","delivery","price",
    )
    points =[]
    for sentence in sentence_split (text ):
        lower =sentence .lower ()
        if any (word in lower for word in signal_words )and 35 <=len (sentence )<=500 :
            points .append (sentence )
        if len (points )>=8 :
            break 

    if not points :
        points =[s for s in sentence_split (text )if 35 <=len (s )<=500 ][:5 ]

    for finding in findings :
        context =finding .get ("context")
        if context and len (points )<8 :
            points .append (context )
    return unique_preserve (points ,8 )


def build_summary (
text :str ,
category :str ,
risk_level :str ,
tags :List [str ],
findings :List [Dict [str ,Any ]],
)->str :
    if not text .strip ():
        return (
        "The file was received, but no machine-readable text was extracted. "
        "OCR may be required for scanned/image-only content."
        )
    risk_count =sum (
    int (item .get ("count")or 0 )
    for item in findings 
    if item .get ("type")=="keyword_match"
    )
    tag_phrase =", ".join (tags [:5 ])if tags else "no specific document tags"
    return (
    f"This {category .lower ()} contains approximately {word_count (text ):,} words. "
    f"The screening engine classified the detected risk level as {risk_level .lower ()} "
    f"and identified {tag_phrase }. The scan found {risk_count } configured "
    "risk-keyword occurrence(s). These signals should be reviewed against "
    "the source document and are not professional advice."
    )


def build_confidence (text :str ,tags :List [str ],entities :List [Dict [str ,Any ]])->int :
    if not text .strip ():
        return 0 
    score =45 
    if len (text )>=500 :
        score +=15 
    if len (text )>=3000 :
        score +=10 
    if tags :
        score +=10 
    if entities :
        score +=10 
    if word_count (text )>=100 :
        score +=5 
    return min (score ,95 )


def run_document_analysis (text :str ,pages :int ,filename :str )->Dict [str ,Any ]:
    text =truncate_text (normalize_text (text ))
    risk_level ,findings =analyze_risk (text )
    tags ,category ,category_description =detect_tags (text )
    entities =extract_entities (text )
    sections =extract_sections (text )
    missing =detect_missing_information (text ,category )
    key_points =build_key_points (text ,findings )
    deadlines =extract_deadlines (text )
    actions =extract_action_items (text )
    obligations =extract_obligations (text )
    sentiment =text_sentiment (text )
    advanced =build_advanced_intelligence (text ,category ,entities ,findings ,deadlines ,obligations )

    return {
    "status":"success",
    "filename":filename ,
    "pages":pages ,
    "word_count":word_count (text ),
    "character_count":len (text ),
    "language":detect_language (text ),
    "risk_level":risk_level ,
    "risk_severity":"Critical" if any(item .get("level")=="Critical" for item in findings )else risk_level ,
    "risk_score":sum (
    int (item .get ("score_contribution")or 0 )
    for item in findings 
    ),
    "risk_summary":{
    "critical":sum (1 for item in findings if item .get ("level")=="Critical"),
    "high":sum (1 for item in findings if item .get ("level")=="High"),
    "medium":sum (1 for item in findings if item .get ("level")=="Medium"),
    "low":sum (1 for item in findings if item .get ("level")=="Low"),
    },
    "confidence":build_confidence (text ,tags ,entities ),
    "tags":tags ,
    "category":category ,
    "category_description":category_description ,
    "summary":build_summary (text ,category ,risk_level ,tags ,findings ),
    "key_points":key_points ,
    "entities":entities ,
    "people":[x ["value"]for x in entities if x .get ("type")=="person"],
    "organizations":[x ["value"]for x in entities if x .get ("type")=="organization"],
    "findings":findings ,
    "risk_explanation":{
    "level":risk_level ,
    "score":min (100 ,sum (int (x .get ("score_contribution")or 0 )for x in findings )),
    "critical_count":sum (1 for x in findings if x .get ("level")=="Critical"),
    "high_count":sum (1 for x in findings if x .get ("level")=="High"),
    "medium_count":sum (1 for x in findings if x .get ("level")=="Medium"),
    "low_count":sum (1 for x in findings if x .get ("level")=="Low"),
    "top_reasons":[x .get ("text")for x in findings [:10 ]],
    "review_priority":"Immediate"if risk_level =="Critical"else "High"if risk_level =="High"else "Normal"
    },
    "missing_information":missing ,
    "sections":sections ,
    "deadlines":deadlines ,
    "action_items":actions ,
    "obligations":obligations ,
    "sentiment":sentiment ,
    "readability":{
    "score":readability_score (text ),
    "level":(
    "Easy"if readability_score (text )>=70 
    else "Moderate"if readability_score (text )>=45 
    else "Difficult"
    ),
    },
    "analysis_engine":"rules-v4-advanced",
    "analyzed_at":utc_now (),
    "advanced_intelligence":advanced ,
    }







CLAUSE_PATTERNS ={
"termination":r"\b(terminat(?:e|ion)|cancel(?:lation)?|end this agreement|notice period)\b",
"confidentiality":r"\b(confidential(?:ity)?|non[- ]disclosure|nda|trade secret)\b",
"indemnity":r"\b(indemnif(?:y|ication)|hold harmless)\b",
"liability":r"\b(liabilit(?:y|ies)|limitation of liability|liable)\b",
"payment":r"\b(payment|invoice|fee|fees|price|pricing|amount due|late fee|interest)\b",
"renewal":r"\b(renew(?:al|s|ed)|auto[- ]renew|automatic renewal)\b",
"intellectual_property":r"\b(intellectual property|copyright|trademark|patent|license|licence|ownership of work)\b",
"data_protection":r"\b(personal data|personal information|privacy|gdpr|data protection|processor|controller)\b",
"warranty":r"\b(warrant(?:y|ies)|warrants|guarantee|guarantees)\b",
"governing_law":r"\b(governing law|jurisdiction|venue|laws of)\b",
"dispute_resolution":r"\b(arbitration|mediation|dispute resolution|court proceedings)\b",
"non_compete":r"\b(non[- ]compete|non[- ]competition|restrictive covenant)\b",
"non_solicit":r"\b(non[- ]solicit|non[- ]solicitation|solicit employees|solicit clients)\b",
"force_majeure":r"\b(force majeure|act of god|unforeseeable event)\b",
"audit":r"\b(audit|auditing|inspection rights|books and records)\b",
"insurance":r"\b(insurance|insured|coverage|certificate of insurance)\b",
}

MONEY_RE =re .compile (
r"(?P<currency>USD|EUR|GBP|PKR|INR|AED|SAR|CAD|AUD|\$|€|£|₨|Rs\.?)\s*"
r"(?P<amount>\d{1,3}(?:[,\s]\d{3})*(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)",re .I )
EMAIL_RE =re .compile (r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",re .I )
PHONE_RE =re .compile (r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)")
URL_RE =re .compile (r"\bhttps?://[^\s<>()]+",re .I )
DATE_RE_ADVANCED =re .compile (
r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|"
r"\d{4}[/-]\d{1,2}[/-]\d{1,2}|"
r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
r"jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|"
r"nov(?:ember)?|dec(?:ember)?)\s+\d{1,2},?\s+\d{4})\b",re .I )

def _dedupe_items (items :List [str ],limit :int =100 )->List [str ]:
    seen ,out =set (),[]
    for item in items :
        clean =normalize_text (str (item ))
        key =clean .lower ()
        if clean and key not in seen :
            seen .add (key )
            out .append (clean )
            if len (out )>=limit :
                break 
    return out 

def extract_advanced_financials (text :str )->List [Dict [str ,Any ]]:
    results =[]
    for m in MONEY_RE .finditer (text ):
        try :
            amount =float (m .group ("amount").replace (",","").replace (" ",""))
        except ValueError :
            amount =None 
        results .append ({
        "value":m .group (0 ).strip (),
        "currency":m .group ("currency"),
        "amount":amount ,
        "context":normalize_text (text [max (0 ,m .start ()-100 ):m .end ()+140 ]),
        })
    return results [:ADVANCED_ENTITY_LIMIT ]

def extract_clause_intelligence (text :str )->List [Dict [str ,Any ]]:
    sentences =sentence_split (text )
    clauses =[]
    for clause_type ,pattern in CLAUSE_PATTERNS .items ():
        rx =re .compile (pattern ,re .I )
        matches =_dedupe_items ([x for x in sentences if rx .search (x )],5 )
        if matches :
            clauses .append ({"type":clause_type ,"count":len (matches ),"passages":matches ,"present":True })
    return clauses 

def extract_risk_explanations (findings :List [Dict [str ,Any ]])->List [Dict [str ,Any ]]:
    out =[]
    for finding in findings [:50 ]:
        level =finding .get ("level","Low")
        title =finding .get ("title")or finding .get ("keyword")or finding .get ("type")or "Risk signal"
        action =(
        "Review immediately and verify with the responsible legal/business owner."
        if level =="Critical"else 
        "Review before approval, signing, payment, or execution."
        if level =="High"else 
        "Confirm the wording, owner, and operational impact."
        if level =="Medium"else 
        "Monitor during normal document review."
        )
        out .append ({"level":level ,"title":title ,
        "reason":finding .get ("description")or finding .get ("context")or title ,
        "recommended_action":action })
    return out 

def extract_sensitive_signals (text :str )->Dict [str ,Any ]:
    emails =_dedupe_items (EMAIL_RE .findall (text ),50 )
    phones =_dedupe_items (PHONE_RE .findall (text ),50 )
    urls =_dedupe_items (URL_RE .findall (text ),50 )
    return {
    "email_count":len (emails ),"phone_count":len (phones ),"url_count":len (urls ),
    "emails":emails ,"phones":phones ,"urls":urls ,
    "contains_personal_contact_data":bool (emails or phones ),
    }

def build_document_health (category ,risk_level ,findings ,missing ,deadlines ,obligations ):
    score =100 
    score -=min (45 ,len (findings )*5 )
    score -=min (30 ,len (missing )*4 )
    score -=10 if risk_level =="Critical"else 7 if risk_level =="High"else 3 if risk_level =="Medium"else 0 
    score =max (0 ,min (100 ,score ))
    return {
    "score":score ,
    "band":"Healthy"if score >=80 else "Review"if score >=55 else "Attention Required",
    "risk_level":risk_level ,"finding_count":len (findings ),
    "missing_information_count":len (missing ),"deadline_count":len (deadlines ),
    "obligation_count":len (obligations ),"category":category ,
    }

def build_advanced_intelligence (text ,category ,entities ,findings ,deadlines ,obligations ):
    missing =detect_missing_information (text ,category )
    risk_level =(
    "Critical"if any (f .get ("level")=="Critical"for f in findings )else 
    "High"if any (f .get ("level")=="High"for f in findings )else 
    "Medium"if any (f .get ("level")=="Medium"for f in findings )else "Low"
    )
    clauses =extract_clause_intelligence (text )
    return {
    "document_health":build_document_health (category ,risk_level ,findings ,missing ,deadlines ,obligations ),
    "clause_intelligence":clauses ,
    "financials":extract_advanced_financials (text ),
    "dates_detected":_dedupe_items (DATE_RE_ADVANCED .findall (text ),100 ),
    "risk_explanations":extract_risk_explanations (findings ),
    "sensitive_data_signals":extract_sensitive_signals (text ),
    "clause_types_detected":[x ["type"]for x in clauses ],
    "entity_count":len (entities ),
    "review_priority":"Immediate"if risk_level =="Critical"else "High"if risk_level =="High"else "Normal",
    }

def build_fast_document_context (record ,hits ):
    analysis =record .get ("analysis")or {}
    payload ={
    "filename":record .get ("filename"),"category":analysis .get ("category"),
    "risk_level":analysis .get ("risk_level"),"risk_score":analysis .get ("risk_score"),
    "summary":analysis .get ("summary"),"key_points":analysis .get ("key_points",[])[:12 ],
    "entities":analysis .get ("entities",[])[:30 ],"deadlines":analysis .get ("deadlines",[])[:30 ],
    "action_items":analysis .get ("action_items",[])[:30 ],"obligations":analysis .get ("obligations",[])[:30 ],
    "missing_information":analysis .get ("missing_information",[])[:30 ],
    "advanced_intelligence":analysis .get ("advanced_intelligence",{}),
    "relevant_passages":[h .get ("text","")for h in hits [:8 ]],
    }
    return truncate_text (json .dumps (payload ,ensure_ascii =False ),MAX_CONTEXT_CHARS )

def fingerprint_text (text ):
    return hashlib .sha256 (text .encode ("utf-8",errors ="ignore")).hexdigest ()

def cached_ai_get (key ):
    with AI_RESPONSE_CACHE_LOCK :
        item =AI_RESPONSE_CACHE .get (key )
        if not item :
            return None 
        created ,value =item 
        if time .time ()-created >AI_CACHE_TTL_SECONDS :
            AI_RESPONSE_CACHE .pop (key ,None )
            return None 
        return value 

def cached_ai_put (key ,value ):
    with AI_RESPONSE_CACHE_LOCK :
        AI_RESPONSE_CACHE [key ]=(time .time (),value )
        if len (AI_RESPONSE_CACHE )>500 :
            oldest =min (AI_RESPONSE_CACHE ,key =lambda k :AI_RESPONSE_CACHE [k ][0 ])
            AI_RESPONSE_CACHE .pop (oldest ,None )





def save_local_record (document_id :str ,filename :str ,analysis :Dict [str ,Any ])->None :
    with LOCAL_LOCK :
        LOCAL_DOCUMENTS [document_id ]={
        "id":document_id ,
        "filename":filename ,
        "analysis":safe_json (analysis ),
        "created_at":utc_now (),
        }


def save_analysis_to_supabase (document_id :str ,filename :str ,analysis :Dict [str ,Any ])->None :
    save_local_record (document_id ,filename ,analysis )
    if not supabase :
        return 
    row ={
    "id":document_id ,
    "filename":filename ,
    "analysis":safe_json (analysis ),
    "created_at":utc_now (),
    }
    try :
        supabase .table ("documents").upsert (row ).execute ()
    except Exception as exc :
        logger .warning ("Supabase save failed: %s",exc )


def get_analysis_from_supabase (document_id :str )->Optional [Dict [str ,Any ]]:
    with LOCAL_LOCK :
        cached =LOCAL_DOCUMENTS .get (document_id )
    if cached :
        return cached 

    if not supabase :
        return None 

    try :
        result =(
        supabase .table ("documents")
        .select ("*")
        .eq ("id",document_id )
        .limit (1 )
        .execute ()
        )
        rows =result .data or []
        if not rows :
            return None 
        with LOCAL_LOCK :
            LOCAL_DOCUMENTS [document_id ]=rows [0 ]
        return rows [0 ]
    except Exception as exc :
        logger .warning ("Supabase lookup failed: %s",exc )
        return None 


def list_records ()->List [Dict [str ,Any ]]:
    with LOCAL_LOCK :
        records =list (LOCAL_DOCUMENTS .values ())
    return sorted (records ,key =lambda x :x .get ("created_at",""),reverse =True )






async def save_upload (file :UploadFile )->Tuple [Path ,str ,int ]:
    original_name =safe_filename (file .filename )
    extension =Path (original_name ).suffix .lower ()

    if extension not in ALLOWED_EXTENSIONS :
        raise HTTPException (
        415 ,
        f"Unsupported file type '{extension or 'unknown'}'. "
        f"Allowed: {', '.join (sorted (ALLOWED_EXTENSIONS ))}",
        )

    document_id =make_document_id (original_name )
    destination =UPLOAD_DIR /f"{document_id }{extension }"
    total_bytes =0 

    try :
        with destination .open ("wb")as buffer :
            while True :
                chunk =await file .read (1024 *1024 )
                if not chunk :
                    break 
                total_bytes +=len (chunk )
                if total_bytes >MAX_FILE_SIZE :
                    raise HTTPException (
                    413 ,
                    f"File exceeds the {MAX_FILE_SIZE_MB } MB maximum upload size.",
                    )
                buffer .write (chunk )
    except HTTPException :
        destination .unlink (missing_ok =True )
        raise 
    except Exception as exc :
        destination .unlink (missing_ok =True )
        raise HTTPException (500 ,f"Could not save file: {exc }")
    finally :
        await file .close ()

    return destination ,original_name ,total_bytes 


def cleanup_old_uploads ()->int :
    cutoff =time .time ()-(FILE_RETENTION_HOURS *3600 )
    removed =0 
    for path in UPLOAD_DIR .iterdir ():
        if not path .is_file ():
            continue 
        try :
            if path .stat ().st_mtime <cutoff :
                path .unlink ()
                removed +=1 
        except OSError :
            continue 
    return removed 






async def call_rapidapi (question :str ,context :str )->Dict [str ,Any ]:
    if not ai_provider_available ():
        raise HTTPException (
        503 ,
        "AI provider is not configured. Set RAPIDAPI_KEY, RAPIDAPI_HOST "
        "and RAPIDAPI_ENDPOINT.",
        )

    payload ={
    "question":question ,
    "context":context [:MAX_CHAT_CONTEXT_CHARS ],
    }
    headers ={
    "x-rapidapi-key":RAPIDAPI_KEY ,
    "x-rapidapi-host":RAPIDAPI_HOST ,
    "Content-Type":"application/json",
    }
    timeout =httpx .Timeout (
    connect =AI_CONNECT_TIMEOUT ,
    read =AI_READ_TIMEOUT ,
    write =10.0 ,
    pool =5.0 ,
    )
    limits =httpx .Limits (max_connections =50 ,max_keepalive_connections =20 )

    try :
        async with httpx .AsyncClient (timeout =timeout ,limits =limits )as client :
            response =await client .post (
            RAPIDAPI_ENDPOINT ,
            headers =headers ,
            json =payload ,
            )
    except httpx .RequestError as exc :
        raise HTTPException (502 ,f"AI provider request failed: {exc }")

    if not response .is_success :
        raise HTTPException (
        502 ,
        f"AI provider returned HTTP {response .status_code }: {response .text [:1000 ]}",
        )

    try :
        provider_data =response .json ()
    except ValueError :
        provider_data ={"raw":response .text }

    return {"provider":"rapidapi","response":provider_data }






def build_search_index (text :str )->List [Dict [str ,Any ]]:
    sentences =sentence_split (text )
    index :List [Dict [str ,Any ]]=[]
    for sentence in sentences [:SEARCH_INDEX_LIMIT ]:
        lower =sentence .lower ()
        words =set (re .findall (r"\w+",lower ))
        index .append ({"text":sentence ,"lower":lower ,"words":words })
    return index 


def search_indexed (index :List [Dict [str ,Any ]],query :str ,limit :int =10 )->List [Dict [str ,Any ]]:
    terms =[t .lower ()for t in re .findall (r"\w+",query )if len (t )>2 ]
    if not terms :
        return []
    scored =[]
    for item in index :
        words =item .get ("words",set ())
        lower =item .get ("lower","")
        overlap =sum (1 for term in terms if term in words )
        frequency =sum (lower .count (term )for term in terms )
        score =overlap *3 +frequency 
        if score :
            scored .append ((score ,item ["text"]))
    scored .sort (key =lambda x :(-x [0 ],len (x [1 ])))
    return [{"score":score ,"text":sentence }for score ,sentence in scored [:limit ]]


def search_text (text :str ,query :str ,limit :int =10 )->List [Dict [str ,Any ]]:
    return search_indexed (build_search_index (text ),query ,limit )


def local_question_answer (question :str ,context :str )->str :
    lower =question .lower ()

    if any (word in lower for word in ("risk","danger","issue","warning")):
        return (
        "The external AI provider is unavailable. Review the extracted "
        "risk_level, findings, and missing_information fields. "
        "Relevant excerpts: "+json .dumps (search_text (context ,question ,5 ),ensure_ascii =False )
        )

    if any (word in lower for word in ("summary","summarize","overview")):
        sentences =sentence_split (context )
        return "Local summary: "+" ".join (sentences [:6 ])

    if any (word in lower for word in ("date","deadline","amount","email","phone","url")):
        entities =extract_entities (context )
        return json .dumps (entities [:30 ],ensure_ascii =False )

    matches =search_text (context ,question ,5 )
    if matches :
        return "Relevant excerpts: "+json .dumps (matches ,ensure_ascii =False )

    return (
    "No external AI provider is configured and no strong local text match "
    "was found for the question."
    )






def build_risk_excel (analysis :Dict [str ,Any ],output_path :Path )->Path :
    """
    Create a professional Excel risk-reminder/report workbook.

    Sheets:
      1. Risk Summary
      2. Risk Findings
      3. Action Items
      4. Deadlines
      5. Missing Information
      6. Entities
    """
    if Workbook is None :
        raise HTTPException (
        501 ,
        "Excel export requires openpyxl. Install it with: pip install openpyxl",
        )

    workbook =Workbook ()
    summary =workbook .active 
    summary .title ="Risk Summary"


    header_fill =PatternFill ("solid",fgColor ="6D28D9")
    critical_fill =PatternFill ("solid",fgColor ="FECACA")
    high_fill =PatternFill ("solid",fgColor ="FED7AA")
    medium_fill =PatternFill ("solid",fgColor ="FEF3C7")
    low_fill =PatternFill ("solid",fgColor ="DCFCE7")
    white_font =Font (color ="FFFFFF",bold =True )
    bold_font =Font (bold =True )
    thin =Side (style ="thin",color ="D1D5DB")
    border =Border (left =thin ,right =thin ,top =thin ,bottom =thin )

    def style_header (ws ,row =1 ):
        for cell in ws [row ]:
            cell .fill =header_fill 
            cell .font =white_font 
            cell .alignment =Alignment (horizontal ="center",vertical ="center")
            cell .border =border 

    def auto_width (ws ):
        for column_cells in ws .columns :
            length =0 
            column_letter =get_column_letter (column_cells [0 ].column )
            for cell in column_cells :
                value =""if cell .value is None else str (cell .value )
                length =max (length ,min (len (value ),70 ))
            ws .column_dimensions [column_letter ].width =max (12 ,length +2 )

    filename =analysis .get ("filename","document")
    risk_level =analysis .get ("risk_level","Unknown")
    risk_score =analysis .get ("risk_score",0 )
    risk_summary =analysis .get ("risk_summary",{})




    summary .append (["DocuSense AI - Risk Report"])
    summary ["A1"].font =Font (size =18 ,bold =True )
    summary .append (["Document",filename ])
    summary .append (["Document ID",analysis .get ("document_id","")])
    summary .append (["Risk Level",risk_level ])
    summary .append (["Risk Score",risk_score ])
    summary .append (["Analyzed At",analysis .get ("analyzed_at","")])
    summary .append (["Category",analysis .get ("category","General document")])
    summary .append (["Confidence",analysis .get ("confidence",0 )])
    summary .append ([])
    summary .append (["Risk Level","Finding Count"])
    for level in ("Critical","High","Medium","Low"):
        summary .append ([level ,risk_summary .get (level .lower (),0 )])

    summary .append ([])
    summary .append (["Executive Summary"])
    summary .append ([analysis .get ("summary","")])

    summary .append ([])
    summary .append (["Recommended Review"])
    if risk_level =="Critical":
        recommendation ="Immediate human review recommended. Prioritize critical clauses before relying on this document."
    elif risk_level =="High":
        recommendation ="Prioritize human review of the high-risk clauses and related obligations."
    elif risk_level =="Medium":
        recommendation ="Review flagged provisions and missing information before finalizing or relying on the document."
    else :
        recommendation ="No major configured risk signal was detected; routine review is still recommended."
    summary .append ([recommendation ])


    summary ["B4"].font =Font (bold =True )
    summary ["B4"].fill ={
    "Critical":critical_fill ,
    "High":high_fill ,
    "Medium":medium_fill ,
    "Low":low_fill ,
    }.get (risk_level ,low_fill )

    for row in summary .iter_rows ():
        for cell in row :
            cell .border =border 
            cell .alignment =Alignment (vertical ="top",wrap_text =True )

    auto_width (summary )
    summary .column_dimensions ["A"].width =28 
    summary .column_dimensions ["B"].width =95 




    findings_ws =workbook .create_sheet ("Risk Findings")
    findings_ws .append ([
    "Priority","Level","Keyword / Signal","Count",
    "Score Contribution","Risk Explanation","Context"
    ])
    style_header (findings_ws )

    for finding in analysis .get ("findings",[]):
        row =findings_ws .max_row +1 
        findings_ws .append ([
        finding .get ("level","Low"),
        finding .get ("level","Low"),
        finding .get ("keyword")or "Document-level signal",
        finding .get ("count",0 ),
        finding .get ("score_contribution",0 ),
        finding .get ("text",""),
        finding .get ("context",""),
        ])
        fill ={
        "Critical":critical_fill ,
        "High":high_fill ,
        "Medium":medium_fill ,
        "Low":low_fill ,
        }.get (finding .get ("level"),low_fill )
        findings_ws .cell (row ,1 ).fill =fill 
        findings_ws .cell (row ,2 ).fill =fill 

    for row in findings_ws .iter_rows ():
        for cell in row :
            cell .border =border 
            cell .alignment =Alignment (vertical ="top",wrap_text =True )
    findings_ws .freeze_panes ="A2"
    auto_width (findings_ws )




    actions_ws =workbook .create_sheet ("Action Items")
    actions_ws .append (["Priority","Action / Obligation","Status","Notes"])
    style_header (actions_ws )

    actions =analysis .get ("action_items",[])
    obligations =analysis .get ("obligations",[])
    combined_actions =unique_preserve (actions +[
    item .get ("text","")for item in obligations 
    if isinstance (item ,dict )
    ],100 )

    for index ,action in enumerate (combined_actions ,1 ):
        priority ="High"if index <=5 else "Medium"
        actions_ws .append ([priority ,action ,"Open","Review against the source document."])

    for row in actions_ws .iter_rows ():
        for cell in row :
            cell .border =border 
            cell .alignment =Alignment (vertical ="top",wrap_text =True )
    actions_ws .freeze_panes ="A2"
    auto_width (actions_ws )




    deadlines_ws =workbook .create_sheet ("Deadlines")
    deadlines_ws .append (["Priority","Deadline / Timing Signal","Status","Reminder"])
    style_header (deadlines_ws )

    for deadline in analysis .get ("deadlines",[]):
        deadlines_ws .append ([
        "High",
        deadline ,
        "Open",
        "Verify the actual deadline in the source document.",
        ])

    for row in deadlines_ws .iter_rows ():
        for cell in row :
            cell .border =border 
            cell .alignment =Alignment (vertical ="top",wrap_text =True )
    deadlines_ws .freeze_panes ="A2"
    auto_width (deadlines_ws )




    missing_ws =workbook .create_sheet ("Missing Information")
    missing_ws .append (["Priority","Missing / Unclear Information","Recommended Check"])
    style_header (missing_ws )

    for item in analysis .get ("missing_information",[]):
        missing_ws .append ([
        "Medium",
        item ,
        "Verify whether this information exists elsewhere in the original document.",
        ])

    for row in missing_ws .iter_rows ():
        for cell in row :
            cell .border =border 
            cell .alignment =Alignment (vertical ="top",wrap_text =True )
    missing_ws .freeze_panes ="A2"
    auto_width (missing_ws )




    entities_ws =workbook .create_sheet ("Entities")
    entities_ws .append (["Type","Label","Value"])
    style_header (entities_ws )

    for entity in analysis .get ("entities",[]):
        entities_ws .append ([
        entity .get ("type",""),
        entity .get ("label",""),
        entity .get ("value",""),
        ])

    for row in entities_ws .iter_rows ():
        for cell in row :
            cell .border =border 
            cell .alignment =Alignment (vertical ="top",wrap_text =True )
    entities_ws .freeze_panes ="A2"
    auto_width (entities_ws )

    workbook .save (output_path )
    return output_path 






@app .get ("/",response_model =HealthResponse )
def home ()->HealthResponse :
    return HealthResponse (
    status ="online",
    service =APP_NAME ,
    version =APP_VERSION ,
    environment =ENVIRONMENT ,
    persistence =persistence_mode (),
    ai_provider ="rapidapi"if ai_provider_available ()else "not-configured",
    )


@app .get ("/health")
def health ()->Dict [str ,Any ]:
    return {
    "status":"healthy",
    "service":APP_NAME ,
    "version":APP_VERSION ,
    "timestamp":utc_now (),
    "documents_cached":len (LOCAL_DOCUMENTS ),
    "checks":{
    "supabase":bool (supabase ),
    "rapidapi":ai_provider_available (),
    "ocr":bool (Image and pytesseract ),
    "docx":bool (DocxDocument ),
    },
    }


@app .get ("/config")
def configuration ()->Dict [str ,Any ]:
    return {
    "max_file_size_mb":MAX_FILE_SIZE_MB ,
    "max_batch_files":MAX_BATCH_FILES ,
    "allowed_extensions":sorted (ALLOWED_EXTENSIONS ),
    "ocr_available":bool (Image and pytesseract ),
    "docx_available":bool (DocxDocument ),
    "persistence_enabled":bool (supabase ),
    "persistence_mode":persistence_mode (),
    "ai_provider_configured":ai_provider_available (),
    "analysis_engine":"rules-v3",
    }


@app .post ("/upload")
async def upload_document (file :UploadFile =File (...))->Dict [str ,Any ]:
    path ,filename ,size =await save_upload (file )
    try :
        pages ,text ,page_data =extract_document (path )
        return {
        "status":"success",
        "document_id":path .stem ,
        "filename":filename ,
        "size_bytes":size ,
        "pages":pages ,
        "characters":len (text ),
        "words":word_count (text ),
        "language":detect_language (text ),
        "preview":text [:5000 ],
        "page_data":page_data ,
        }
    finally :
        path .unlink (missing_ok =True )


@app .post ("/analyze")
async def analyze_document (file :UploadFile =File (...))->Dict [str ,Any ]:
    started =time .perf_counter ()
    path ,filename ,size =await save_upload (file )
    document_id =path .stem 

    try :
        pages ,text ,page_data =extract_document (path )
        analysis =run_document_analysis (text ,pages ,filename )
        analysis .update ({
        "document_id":document_id ,
        "size_bytes":size ,
        "processing_ms":round ((time .perf_counter ()-started )*1000 ,2 ),
        "page_data":page_data ,
        })
        save_analysis_to_supabase (document_id ,filename ,analysis )
        return analysis 
    finally :
        path .unlink (missing_ok =True )


@app .post ("/analyze/{document_id}/risk-report")
def generate_risk_report (document_id :str ):
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document analysis not found.")

    analysis =record .get ("analysis")or {}
    filename =safe_filename (record .get ("filename","document"))
    report_dir =UPLOAD_DIR /"reports"
    report_dir .mkdir (parents =True ,exist_ok =True )

    output_path =report_dir /f"{Path (filename ).stem }_risk_report_{document_id }.xlsx"
    build_risk_excel (analysis ,output_path )

    return {
    "status":"success",
    "document_id":document_id ,
    "filename":output_path .name ,
    "path":str (output_path ),
    "message":"Excel risk report generated successfully.",
    }


@app .post ("/analyze/batch")
async def analyze_batch (files :List [UploadFile ]=File (...))->Dict [str ,Any ]:
    if not files :
        raise HTTPException (400 ,"No files supplied.")
    if len (files )>MAX_BATCH_FILES :
        raise HTTPException (413 ,f"Maximum batch size is {MAX_BATCH_FILES } files.")

    semaphore =asyncio .Semaphore (max (1 ,BATCH_CONCURRENCY ))

    async def process_one (upload :UploadFile )->Dict [str ,Any ]:
        async with semaphore :
            try :
                result =await analyze_document (upload )
                return {
                "status":"success",
                "filename":upload .filename ,
                "analysis":result .get ("analysis"),
                }
            except HTTPException as exc :
                return {
                "status":"error",
                "filename":upload .filename ,
                "error":exc .detail ,
                }
            except Exception as exc :
                logger .exception ("Batch analysis failed for %s",upload .filename )
                return {
                "status":"error",
                "filename":upload .filename ,
                "error":str (exc ),
                }

    results =await asyncio .gather (*(process_one (upload )for upload in files ))
    return {
    "status":"success",
    "count":len (results ),
    "results":results ,
    }


@app .get ("/documents")
def documents (
category :Optional [str ]=Query (default =None ),
risk_level :Optional [str ]=Query (default =None ),
q :Optional [str ]=Query (default =None ,max_length =200 ),
limit :int =Query (default =50 ,ge =1 ,le =200 ),
)->Dict [str ,Any ]:
    records =list_records ()
    filtered =[]

    for record in records :
        analysis =record .get ("analysis")or {}
        if category and analysis .get ("category","").lower ()!=category .lower ():
            continue 
        if risk_level and analysis .get ("risk_level","").lower ()!=risk_level .lower ():
            continue 
        if q :
            haystack =(
            f"{record .get ('filename','')} "
            f"{analysis .get ('summary','')} "
            f"{' '.join (analysis .get ('tags',[]))}"
            ).lower ()
            if q .lower ()not in haystack :
                continue 

        filtered .append ({
        "document_id":record .get ("id"),
        "filename":record .get ("filename"),
        "created_at":record .get ("created_at"),
        "category":analysis .get ("category"),
        "risk_level":analysis .get ("risk_level"),
        "confidence":analysis .get ("confidence"),
        "pages":analysis .get ("pages"),
        "word_count":analysis .get ("word_count"),
        "tags":analysis .get ("tags",[]),
        })

    return {
    "status":"success",
    "count":len (filtered [:limit ]),
    "documents":filtered [:limit ],
    }


@app .get ("/documents/{document_id}")
def get_document (document_id :str )->Dict [str ,Any ]:
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document analysis not found.")
    return {"status":"success","document":record }


@app .post ("/documents/{document_id}/reanalyze")
async def reanalyze_document (document_id :str )->Dict [str ,Any ]:
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document analysis not found.")

    analysis =record .get ("analysis")or {}

    page_data =analysis .get ("page_data")or []
    text ="\n".join (
    item .get ("text","")for item in page_data if isinstance (item ,dict )
    )

    if not text .strip ():
        raise HTTPException (
        422 ,
        "Original extracted text is not available for re-analysis. "
        "Upload the source document again.",
        )

    refreshed =await run_in_threadpool (
    run_document_analysis ,
    text ,
    analysis .get ("pages",1 ),
    record .get ("filename","document"),
    )
    refreshed .update ({
    "document_id":document_id ,
    "size_bytes":analysis .get ("size_bytes",0 ),
    "page_data":page_data ,
    })
    await run_in_threadpool (
    save_analysis_to_supabase ,
    document_id ,
    record .get ("filename","document"),
    refreshed ,
    )
    cache_document (document_id ,{
    "id":document_id ,
    "filename":record .get ("filename","document"),
    "analysis":refreshed ,
    "created_at":utc_now (),
    })
    return {"status":"success","analysis":refreshed }


@app .get ("/documents/{document_id}/search")
def search_document (
document_id :str ,
q :str =Query (...,min_length =1 ,max_length =500 ),
limit :int =Query (default =10 ,ge =1 ,le =50 ),
)->Dict [str ,Any ]:
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")

    cached =get_cached_document (document_id )
    if not cached :
        cache_document (document_id ,record )
        cached =get_cached_document (document_id )
    results =search_indexed ((cached or {}).get ("search_index",[]),q ,limit )
    return {
    "status":"success",
    "document_id":document_id ,
    "query":q ,
    "results":results ,
    }


@app .get ("/documents/{document_id}/stats")
def document_stats (document_id :str )->Dict [str ,Any ]:
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")

    analysis =record .get ("analysis")or {}
    findings =analysis .get ("findings",[])
    return {
    "status":"success",
    "document_id":document_id ,
    "filename":record .get ("filename"),
    "statistics":{
    "pages":analysis .get ("pages",0 ),
    "words":analysis .get ("word_count",0 ),
    "characters":analysis .get ("character_count",0 ),
    "sections":len (analysis .get ("sections",[])),
    "entities":len (analysis .get ("entities",[])),
    "risk_findings":sum (
    int (item .get ("count")or 0 )
    for item in findings 
    if item .get ("type")=="keyword_match"
    ),
    "deadlines":len (analysis .get ("deadlines",[])),
    "action_items":len (analysis .get ("action_items",[])),
    "obligations":len (analysis .get ("obligations",[])),
    "missing_information":len (analysis .get ("missing_information",[])),
    "readability":analysis .get ("readability",{}),
    "sentiment":analysis .get ("sentiment",{}),
    },
    }


@app .post ("/compare")
def compare_documents (request :CompareRequest )->Dict [str ,Any ]:
    documents =[]
    for document_id in request .document_ids :
        record =get_analysis_from_supabase (document_id )
        if record :
            documents .append (record )

    if len (documents )<2 :
        raise HTTPException (
        404 ,
        "At least two requested documents must exist to compare them.",
        )

    analyses =[document .get ("analysis")or {}for document in documents ]
    categories =Counter (a .get ("category",DEFAULT_CATEGORY )for a in analyses )
    risks =Counter (a .get ("risk_level","Unknown")for a in analyses )

    all_tags =unique_preserve (
    [tag for a in analyses for tag in a .get ("tags",[])],50 
    )

    tag_sets =[set (a .get ("tags",[]))for a in analyses ]
    shared_tags =sorted (set .intersection (*tag_sets ))if tag_sets else []

    return {
    "status":"success",
    "document_count":len (documents ),
    "documents":[
    {
    "document_id":document .get ("id"),
    "filename":document .get ("filename"),
    "category":(document .get ("analysis")or {}).get ("category"),
    "risk_level":(document .get ("analysis")or {}).get ("risk_level"),
    "confidence":(document .get ("analysis")or {}).get ("confidence"),
    "word_count":(document .get ("analysis")or {}).get ("word_count"),
    "tags":(document .get ("analysis")or {}).get ("tags",[]),
    }
    for document in documents 
    ],
    "comparison":{
    "categories":dict (categories ),
    "risk_levels":dict (risks ),
    "all_tags":all_tags ,
    "shared_tags":shared_tags ,
    },
    }


@app .get ("/search")
def global_search (
q :str =Query (...,min_length =1 ,max_length =500 ),
limit :int =Query (default =20 ,ge =1 ,le =100 ),
)->Dict [str ,Any ]:
    results =[]
    query_terms =[t .lower ()for t in re .findall (r"\w+",q )if len (t )>2 ]

    for record in list_records ():
        analysis =record .get ("analysis")or {}
        searchable =" ".join ([
        record .get ("filename",""),
        analysis .get ("summary",""),
        " ".join (analysis .get ("tags",[])),
        " ".join (analysis .get ("key_points",[])),
        ]).lower ()

        score =sum (searchable .count (term )for term in query_terms )
        if score :
            results .append ({
            "document_id":record .get ("id"),
            "filename":record .get ("filename"),
            "score":score ,
            "category":analysis .get ("category"),
            "risk_level":analysis .get ("risk_level"),
            })

    results .sort (key =lambda item :(-item ["score"],item ["filename"]))
    return {"status":"success","query":q ,"results":results [:limit ]}


@app .get ("/analytics")
def analytics ()->Dict [str ,Any ]:
    records =list_records ()
    analyses =[r .get ("analysis")or {}for r in records ]

    categories =Counter (a .get ("category",DEFAULT_CATEGORY )for a in analyses )
    risks =Counter (a .get ("risk_level","Unknown")for a in analyses )
    languages =Counter (a .get ("language","Unknown")for a in analyses )

    total_words =sum (int (a .get ("word_count")or 0 )for a in analyses )
    total_pages =sum (int (a .get ("pages")or 0 )for a in analyses )
    total_findings =sum (
    sum (int (x .get ("count")or 0 )for x in a .get ("findings",[]))
    for a in analyses 
    )

    return {
    "status":"success",
    "document_count":len (records ),
    "total_words":total_words ,
    "total_pages":total_pages ,
    "total_risk_signals":total_findings ,
    "categories":dict (categories ),
    "risk_levels":dict (risks ),
    "languages":dict (languages ),
    "recent_documents":[
    {
    "document_id":r .get ("id"),
    "filename":r .get ("filename"),
    "created_at":r .get ("created_at"),
    "category":(r .get ("analysis")or {}).get ("category"),
    "risk_level":(r .get ("analysis")or {}).get ("risk_level"),
    }
    for r in records [:10 ]
    ],
    }


@app .post ("/chat")
async def chat (request :ChatRequest )->Dict [str ,Any ]:
    question =request .question .strip ()
    context =request .context or ""
    search_hits :List [Dict [str ,Any ]]=[]
    analysis :Dict [str ,Any ]={}
    record :Optional [Dict [str ,Any ]]=None 

    if request .document_id :
        cached =get_cached_document (request .document_id )
        if cached :
            record =cached ["record"]
            analysis =record .get ("analysis")or {}
            search_hits =search_indexed (cached .get ("search_index",[]),question ,6 )
        else :
            record =get_analysis_from_supabase (request .document_id )
            if not record :
                raise HTTPException (404 ,"Document not found.")
            analysis =record .get ("analysis")or {}
            page_data =analysis .get ("page_data")or []
            text ="\n".join (
            item .get ("text","")for item in page_data if isinstance (item ,dict )
            )
            cache_document (request .document_id ,record )
            cached =get_cached_document (request .document_id )
            search_hits =search_indexed ((cached or {}).get ("search_index",[]),question ,6 )



        q =question .lower ()
        if any (k in q for k in ("risk","danger","warning","issue")):
            return {
            "status":"success","question":question ,
            "document_id":request .document_id ,"mode":"local-fast",
            "response":json .dumps ({
            "risk_level":analysis .get ("risk_level"),
            "risk_score":analysis .get ("risk_score"),
            "risk_summary":analysis .get ("risk_summary"),
            "findings":analysis .get ("findings",[])[:10 ],
            },ensure_ascii =False ),
            }
        if any (k in q for k in ("summary","summarize","overview")):
            return {
            "status":"success","question":question ,
            "document_id":request .document_id ,"mode":"local-fast",
            "response":analysis .get ("summary","No summary available."),
            }
        if any (k in q for k in ("deadline","due date","expires","expiration","renewal")):
            return {
            "status":"success","question":question ,
            "document_id":request .document_id ,"mode":"local-fast",
            "response":json .dumps ({
            "deadlines":analysis .get ("deadlines",[]),
            },ensure_ascii =False ),
            }

        context =json .dumps ({
        "filename":record .get ("filename")if record else None ,
        "summary":analysis .get ("summary"),
        "key_points":analysis .get ("key_points"),
        "findings":analysis .get ("findings"),
        "entities":analysis .get ("entities"),
        "missing_information":analysis .get ("missing_information"),
        "sections":analysis .get ("sections"),
        "deadlines":analysis .get ("deadlines"),
        "action_items":analysis .get ("action_items"),
        "obligations":analysis .get ("obligations"),
        "relevant_passages":[hit ["text"]for hit in search_hits ],
        },ensure_ascii =False )

    if not context :
        raise HTTPException (422 ,"Provide document_id or context so the question can be answered.")

    if ai_provider_available ():
        ai_result =await call_rapidapi (question ,context )
        return {
        "status":"success","question":question ,
        "document_id":request .document_id ,"mode":"ai",**ai_result ,
        }

    return {
    "status":"success","question":question ,
    "document_id":request .document_id ,"mode":"local",
    "response":local_question_answer (question ,context ),
    "note":"External AI is not configured; this answer uses local document signals.",
    }


@app .get ("/metrics")
def metrics ()->Dict [str ,Any ]:
    records =list_records ()
    risk_counter =Counter ()
    category_counter =Counter ()
    for record in records :
        analysis =record .get ("analysis")or {}
        risk_counter [analysis .get ("risk_level","Unknown")]+=1 
        category_counter [analysis .get ("category",DEFAULT_CATEGORY )]+=1 

    return {
    "status":"success",
    "uptime_note":"Process-local metrics",
    "documents_cached":len (records ),
    "risk_distribution":dict (risk_counter ),
    "category_distribution":dict (category_counter ),
    "rate_limit_per_minute":RATE_LIMIT_PER_MINUTE ,
    }


@app .exception_handler (HTTPException )
async def http_exception_handler (request :Request ,exc :HTTPException ):
    request_id =request .headers .get ("X-Request-ID","unknown")
    return JSONResponse (
    status_code =exc .status_code ,
    content ={
    "status":"error",
    "detail":exc .detail ,
    "request_id":request_id ,
    },
    )


@app .exception_handler (Exception )
async def unhandled_exception_handler (request :Request ,exc :Exception ):
    request_id =request .headers .get ("X-Request-ID","unknown")
    logger .exception ("Unhandled exception | request_id=%s",request_id )
    return JSONResponse (
    status_code =500 ,
    content ={
    "status":"error",
    "detail":"Internal server error.",
    "request_id":request_id ,
    },
    )







@app .get ("/documents/{document_id}/intelligence")
def document_intelligence (document_id :str ):
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")
    return {"status":"success","document_id":document_id ,
    "filename":record .get ("filename"),
    "intelligence":(record .get ("analysis")or {}).get ("advanced_intelligence",{})}

@app .get ("/documents/{document_id}/timeline")
def document_timeline (document_id :str ):
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")
    analysis =record .get ("analysis")or {}
    events =[]
    for item in analysis .get ("deadlines",[]):
        if isinstance (item ,dict ):
            events .append ({"type":"deadline","date":item .get ("date")or item .get ("value"),
            "title":item .get ("label")or item .get ("type")or "Important date","details":item })
        else :
            events .append ({"type":"deadline","date":str (item ),"title":"Important date"})
    for date_value in ((analysis .get ("advanced_intelligence")or {}).get ("dates_detected",[])):
        events .append ({"type":"detected_date","date":date_value ,"title":"Detected date"})
    return {"status":"success","document_id":document_id ,"timeline":events }

@app .get ("/documents/{document_id}/clauses")
def document_clauses (document_id :str ,clause_type :Optional [str ]=Query (default =None ,max_length =100 )):
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")
    clauses =((record .get ("analysis")or {}).get ("advanced_intelligence")or {}).get ("clause_intelligence",[])
    if clause_type :
        clauses =[c for c in clauses if c .get ("type","").lower ()==clause_type .lower ()]
    return {"status":"success","document_id":document_id ,"clauses":clauses }

@app .get ("/documents/{document_id}/financials")
def document_financials (document_id :str ):
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")
    financials =((record .get ("analysis")or {}).get ("advanced_intelligence")or {}).get ("financials",[])
    totals =defaultdict (float )
    for item in financials :
        if isinstance (item .get ("amount"),(int ,float )):
            totals [str (item .get ("currency")or "UNKNOWN").upper ()]+=item ["amount"]
    return {"status":"success","document_id":document_id ,
    "financial_mentions":financials ,"totals_by_currency":dict (totals )}

@app .get ("/documents/{document_id}/review")
def document_review (document_id :str ):
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")
    analysis =record .get ("analysis")or {}
    advanced =analysis .get ("advanced_intelligence")or {}
    recommendations =[x .get ("recommended_action")for x in advanced .get ("risk_explanations",[])]
    recommendations +=[f"Verify missing information: {x }"for x in analysis .get ("missing_information",[])]
    return {
    "status":"success","document_id":document_id ,
    "review_priority":advanced .get ("review_priority"),
    "health":advanced .get ("document_health"),
    "critical_findings":[f for f in analysis .get ("findings",[])if f .get ("level")in ("Critical","High")],
    "recommendations":_dedupe_items (recommendations ,30 ),
    }

@app .post ("/documents/{document_id}/ask")
async def ask_document (request :ChatRequest ,document_id :str ):
    request .document_id =document_id 
    return await chat (request )

@app .post ("/ai/enrich/{document_id}")
async def enrich_document (document_id :str ):
    record =get_analysis_from_supabase (document_id )
    if not record :
        raise HTTPException (404 ,"Document not found.")
    cached =get_cached_document (document_id )
    hits =search_indexed ((cached or {}).get ("search_index",[]),"summary risks obligations deadlines clauses",8 )
    context =build_fast_document_context (record ,hits )
    if not ai_provider_available ():
        return {"status":"success","mode":"local",
        "message":"AI provider is not configured.","context":context }
    analysis =record .get ("analysis")or {}
    key =fingerprint_text (f"enrich:{document_id }:{analysis .get ('analyzed_at')}")
    cached_result =cached_ai_get (key )
    if cached_result :
        return {"status":"success","mode":"cache",**cached_result }
    result =await call_rapidapi (
    "Analyze this document for executive insights, risks, obligations, important clauses, missing information and recommended next actions. Prefer structured JSON.",
    context )
    cached_ai_put (key ,result )
    return {"status":"success","mode":"ai",**result }

@app .get ("/system/capabilities")
def capabilities ():
    return {"status":"success","version":APP_VERSION ,"features":{
    "pdf":True ,"docx":DocxDocument is not None ,
    "ocr":pytesseract is not None and Image is not None ,
    "supabase":bool (supabase ),"ai_provider":ai_provider_available (),
    "advanced_clause_intelligence":True ,"financial_extraction":True ,
    "document_health_scoring":True ,"indexed_search":True ,
    "fast_local_qa":True ,"batch_concurrency":BATCH_CONCURRENCY ,
    }}





if __name__ =="__main__":
    import uvicorn 

    uvicorn .run (
    "main:app",
    host =os .getenv ("HOST","127.0.0.1"),
    port =int (os .getenv ("PORT","8000")),
    reload =ENVIRONMENT =="development",
    )
