from __future__ import annotations

import csv
import io
import logging
import os
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from pypdf import PdfReader

try:
    import fitz
except ImportError:
    fitz = None
try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None
try:
    from PIL import Image
except ImportError:
    Image = None
try:
    import pytesseract
except ImportError:
    pytesseract = None

from config import Settings
from utils.text import normalize_text, word_count

logger = logging.getLogger("docusense.document")

class DocumentProcessingError(RuntimeError): pass

class DocumentService:
    def __init__(self, settings: Settings):
        self.settings = settings
        if settings.tesseract_path and pytesseract:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_path

    def extract(self, path: Path) -> tuple[int, str, list[dict[str, Any]]]:
        ext = path.suffix.lower()
        if ext == ".pdf": return self.extract_pdf(path)
        if ext == ".docx": return self.extract_docx(path)
        if ext in {".txt", ".md"}: return self.extract_text(path)
        if ext == ".csv": return self.extract_csv(path)
        if ext in {".png", ".jpg", ".jpeg", ".webp"}: return self.extract_image(path)
        raise HTTPException(415, f"Unsupported file type: {ext or 'unknown'}")

    def extract_pdf(self, path: Path) -> tuple[int, str, list[dict[str, Any]]]:
        if fitz:
            try:
                return self._extract_pdf_fitz(path)
            except Exception as exc:
                logger.warning("PyMuPDF extraction failed; using pypdf fallback: %s", exc)
        try:
            reader = PdfReader(str(path))
        except Exception as exc:
            raise HTTPException(400, "Could not read PDF file.") from exc
        if len(reader.pages) > self.settings.max_document_pages:
            raise HTTPException(413, f"PDF exceeds the {self.settings.max_document_pages}-page limit.")
        pages=[]; missing=[]
        for number,page in enumerate(reader.pages,1):
            try: text=normalize_text(page.extract_text() or "")
            except Exception: text=""
            pages.append({"page":number,"text":text,"characters":len(text),"words":word_count(text),"ocr":False})
            if len(text.strip()) < 20: missing.append(number)
        self._ocr_missing_pdf_pages(path,pages,missing)
        return len(pages), normalize_text("\n\n".join(p["text"] for p in pages if p["text"])), pages

    def _extract_pdf_fitz(self, path: Path):
        with fitz.open(str(path)) as doc:
            if len(doc) > self.settings.max_document_pages:
                raise HTTPException(413, f"PDF exceeds the {self.settings.max_document_pages}-page limit.")
            pages=[]; missing=[]
            for i,page in enumerate(doc):
                text=normalize_text(page.get_text("text", sort=True) or "")
                item={"page":i+1,"text":text,"characters":len(text),"words":word_count(text),"ocr":False}
                pages.append(item)
                if len(text.strip()) < 20: missing.append(i+1)
        self._ocr_missing_pdf_pages(path,pages,missing)
        return len(pages), normalize_text("\n\n".join(p["text"] for p in pages if p["text"])), pages

    def _ocr_missing_pdf_pages(self,path:Path,pages:list[dict[str,Any]],missing:list[int]) -> None:
        if not missing or not self.settings.ocr_enabled or not (fitz and pytesseract and Image): return
        def work(number:int):
            try:
                with fitz.open(str(path)) as doc:
                    page=doc.load_page(number-1)
                    pix=page.get_pixmap(matrix=fitz.Matrix(1.45,1.45),alpha=False,colorspace=fitz.csRGB)
                    image=Image.frombytes("RGB",[pix.width,pix.height],pix.samples)
                text=normalize_text(pytesseract.image_to_string(image,config="--oem 1 --psm 6"))
                return number,text
            except Exception as exc:
                logger.warning("OCR failed for page %s: %s",number,exc); return number,""
        with ThreadPoolExecutor(max_workers=max(1,min(self.settings.ocr_workers,len(missing)))) as pool:
            for number,text in pool.map(work,missing):
                page=pages[number-1]; page["text"]=text; page["characters"]=len(text); page["words"]=word_count(text); page["ocr"]=bool(text)

    def extract_docx(self,path:Path):
        if not DocxDocument: raise HTTPException(501,"DOCX support requires python-docx.")
        try: document=DocxDocument(str(path))
        except Exception as exc: raise HTTPException(400,"Could not read DOCX file.") from exc
        parts=[]
        for paragraph in document.paragraphs:
            text=normalize_text(paragraph.text)
            if text: parts.append(text)
        for index,table in enumerate(document.tables,1):
            parts.append(f"[Table {index}]")
            for row in table.rows: parts.append(" | ".join(normalize_text(cell.text) for cell in row.cells))
        text=normalize_text("\n".join(parts))
        return 1,text,[{"page":1,"text":text,"characters":len(text),"words":word_count(text),"tables":len(document.tables)}]

    def extract_text(self,path:Path):
        try: text=normalize_text(path.read_text(encoding="utf-8",errors="replace"))
        except Exception as exc: raise HTTPException(400,"Could not read text file.") from exc
        return 1,text,[{"page":1,"text":text,"characters":len(text),"words":word_count(text)}]

    def extract_csv(self,path:Path):
        try: raw=path.read_text(encoding="utf-8",errors="replace"); rows=list(csv.reader(io.StringIO(raw)))
        except Exception as exc: raise HTTPException(400,"Could not read CSV file.") from exc
        text=normalize_text("\n".join(" | ".join(c.strip() for c in row) for row in rows))
        return 1,text,[{"page":1,"text":text,"characters":len(text),"words":word_count(text),"rows":len(rows),"columns":max((len(r) for r in rows),default=0)}]

    def extract_image(self,path:Path):
        if not Image or not pytesseract: raise HTTPException(501,"Image OCR requires Pillow and pytesseract.")
        try: pytesseract.get_tesseract_version()
        except Exception as exc: raise HTTPException(501,"Tesseract OCR is not available. Configure TESSERACT_PATH.") from exc
        try:
            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image: text=normalize_text(pytesseract.image_to_string(image))
        except HTTPException: raise
        except Exception as exc: raise HTTPException(400,"Could not process image.") from exc
        return 1,text,[{"page":1,"text":text,"characters":len(text),"words":word_count(text),"ocr":True}]
