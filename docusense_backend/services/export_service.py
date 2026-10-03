from __future__ import annotations
from pathlib import Path
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExportService:
    def risk_report(self,analysis:dict[str,Any],path:Path)->Path:
        wb=Workbook(); ws=wb.active; ws.title="Risk Summary"
        header=PatternFill("solid",fgColor="111827"); white=Font(color="FFFFFF",bold=True); thin=Side(style="thin",color="D1D5DB")
        border=Border(left=thin,right=thin,top=thin,bottom=thin)
        rows=[["Document",analysis.get("filename","")],["Document ID",analysis.get("document_id","")],["Risk Level",analysis.get("risk_level","")],["Risk Score",analysis.get("risk_score",0)],["Category",analysis.get("category","")],["Confidence",analysis.get("confidence",0)],["Executive Summary",analysis.get("executive_summary","")]]
        for row in rows: ws.append(row)
        for cell in ws[1]: cell.font=white; cell.fill=header
        for row in ws.iter_rows():
            for cell in row: cell.border=border; cell.alignment=Alignment(vertical="top",wrap_text=True)
        self._simple_sheet(wb,"Risk Findings",["Level","Signal","Count","Explanation","Page","Evidence"],[[x.get("level"),x.get("keyword"),x.get("count"),x.get("text"),x.get("page"),x.get("evidence") or x.get("context")] for x in analysis.get("findings",[])])
        self._simple_sheet(wb,"Action Items",["Action","Type"],[[x,"action"] for x in analysis.get("action_items",[])])
        self._simple_sheet(wb,"Deadlines",["Date / Timing","Page","Evidence"],[[x.get("date") if isinstance(x,dict) else x,x.get("page") if isinstance(x,dict) else None,x.get("evidence","") if isinstance(x,dict) else ""] for x in analysis.get("timeline",analysis.get("deadlines",[]))])
        self._simple_sheet(wb,"Missing Information",["Item"],[[x] for x in analysis.get("missing_information",[])])
        self._simple_sheet(wb,"Entities",["Type","Value","Page","Evidence"],[[x.get("type"),x.get("value"),x.get("page"),x.get("evidence")] for x in analysis.get("entities",[])])
        path.parent.mkdir(parents=True,exist_ok=True); wb.save(path); return path

    def cv_batch(self,rows:list[dict[str,Any]],path:Path)->Path:
        wb=Workbook(); ws=wb.active; ws.title="CV Analysis"
        fields=["filename","name","email","phone","location","education","skills","experience","companies","job_titles","years_experience","certifications","languages"]
        ws.append(fields)
        for row in rows: ws.append([row.get(x,"") for x in fields])
        self._format(ws); path.parent.mkdir(parents=True,exist_ok=True); wb.save(path); return path

    def _simple_sheet(self,wb,name,headers,rows):
        ws=wb.create_sheet(name); ws.append(headers)
        for row in rows: ws.append(row)
        self._format(ws)
    def _format(self,ws):
        fill=PatternFill("solid",fgColor="111827"); font=Font(color="FFFFFF",bold=True); thin=Side(style="thin",color="D1D5DB")
        for cell in ws[1]: cell.fill=fill; cell.font=font
        for row in ws.iter_rows():
            for cell in row: cell.border=Border(left=thin,right=thin,top=thin,bottom=thin); cell.alignment=Alignment(vertical="top",wrap_text=True)
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for cells in ws.columns:
            letter=get_column_letter(cells[0].column); width=min(70,max(12,max(len(str(c.value or "")) for c in cells)+2)); ws.column_dimensions[letter].width=width
